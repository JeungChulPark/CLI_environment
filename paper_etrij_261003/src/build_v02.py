"""Manuscript v0.2 (English, ETRI Journal format), restructured to follow the overview figure
(images/fig3_real_v2.png):

  A. Asynchronous perception and capture-time registration
     1 sensor platform and time base, SLAM stream, 2 zero-shot 6D estimation, 3 asynchronous timing,
     4 capture-time registration  -> anchored object observation
  B. Keyframe-anchored spatial object mapping
     5 keyframe-anchored object state map, 6 loop-closure-aware update
  7. Result

Formatting helpers and references come from build_docx.py.
"""
import os
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import build_docx as B
from build_docx import W, runs, fmt, set_cols, new_doc, add_page_numbers, COL_W, FULL_W, REFS

TITLE = ('Keyframe-Anchored Object Maps: Registering Late Zero-Shot 6D Pose Estimates '
         'in Asynchronous SLAM')
ABSTRACT = (
    "A zero-shot 6D pose estimator recognizes objects from CAD models alone but needs about a second per image. Running it "
    "on a second computer keeps SLAM real-time, yet its results arrive 0.7–6.6 s after capture, from another camera and "
    "clock, after loop closure may have corrected the map. This study presents a two-stage method that maps such late "
    "results consistently. In the first stage, asynchronous perception and capture-time registration, every result is "
    "registered with the SLAM pose interpolated at its capture time and expressed relative to the reference keyframe of "
    "that time. In the second stage, a keyframe-anchored object state map stores each object as an identifier, anchor "
    "keyframe, relative pose, existence probability, and class; a loop closure updates only keyframe poses, so every "
    "object follows the correction. On 61 replayed live runs, the median deviation of repeated observations drops from "
    "9.67 cm (arrival-time pose) to 1.51 cm, and when the map changes in flight the share above 10 cm drops from 30.8% to "
    "1.5%. Indoors, 72 of 74 reappearing objects are re-associated.")
KEYWORDS = ('asynchronous sensor fusion, loop closure, object-level mapping, '
            'simultaneous localization and mapping, zero-shot 6D object pose estimation')


def front(d, title, abstract, keywords, labels=('ORIGINAL ARTICLE', 'Abstract', 'KEYWORDS'), title_size=17):
    set_cols(d.sections[0], 1)
    p = d.add_paragraph(); r = p.add_run(labels[0]); r.bold = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x1F, 0x5F, 0xA8); fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, indent=False)
    p = d.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(title_size)
    fmt(p, size=title_size, align=WD_ALIGN_PARAGRAPH.LEFT, after=14, indent=False, line=1.05)
    for head, body, after in ((labels[1], abstract, 6), (labels[2], keywords, 12)):
        p = d.add_paragraph(); r = p.add_run(head); r.bold = True; r.font.size = Pt(9.5)
        fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=2, indent=False)
        p = runs(d.add_paragraph(), body, size=9.5); fmt(p, size=9.5, after=after, indent=False)
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '6'), ('space', '6'), ('color', '000000')):
        b.set(qn(f'w:{k}'), v)
    bd.append(b); pPr.append(bd)
    s = d.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 2)


def refs(d, head='REFERENCES'):
    p = d.add_paragraph(); r = p.add_run(head); r.bold = True; r.font.size = Pt(10)
    fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=4, indent=False)
    for i, ref in enumerate(REFS, 1):
        p = d.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
        r = p.add_run(f'{i}.\t'); r.font.size = Pt(8); runs(p, ref, size=8)
        fmt(p, size=8, align=WD_ALIGN_PARAGRAPH.LEFT, after=1.5, indent=False)
        p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.first_line_indent = Cm(-0.7)


S = Cm


def build(path):
    assert len(ABSTRACT) <= 1200, len(ABSTRACT)
    d = new_doc(); w = W(d)
    front(d, TITLE, ABSTRACT, KEYWORDS)

    # ================================================================ 1 Introduction
    w.H1(1, 'Introduction')
    w.P("Service and logistics robots, digital twins, and XR applications need *object-level* maps: the objects in a space, "
        "each with a 3D position and orientation in the map frame, kept up to date while the platform moves. Visual and "
        "LiDAR SLAM systems such as ORB-SLAM3 [1] and HDL-Graph-SLAM [2] estimate the sensor pose in real time and remove "
        "accumulated drift by loop closure. Zero-shot 6D object pose estimators such as SAM-6D [3], FoundPose [4], and "
        "GigaPose [5] estimate the pose of an object unseen in training from its CAD model alone.")
    w.P("The two do not combine by multiplying two transforms. Zero-shot estimators are slow, and an estimator that shares "
        "a computer with SLAM takes its time budget: in our system, SLAM pose delivery latency rose from 0.04 s to 2.5 s when "
        "a heavy process ran on the same computer, and a slower SLAM tracks less reliably and corrects its map more often. "
        "Fusion++ [6], which runs recognition and SLAM together, operates at 4–8 Hz overall. Running recognition on a second "
        "computer keeps SLAM real-time, but then each result arrives seconds after its image was captured, from a different "
        "camera and clock, and SLAM may correct its past keyframes by loop closure while the result is in flight.")
    w.P("Interpolating the trajectory at the capture timestamp, as a timestamped transform buffer such as ROS tf [7] does, "
        "places a late result correctly at the time of registration, but it stores the object in the map frame; after a "
        "loop closure moves the trajectory, the object stays behind. Figure 1 (panel 6) shows a real case: after a loop "
        "closure that moved keyframes by up to 3.05 m, 68 of 119 results stored this way lie more than 10 cm from their final "
        "object position, versus 6 of 119 with the method proposed here.")
    w.P("This study organizes the solution in two stages (Fig. 1). The contributions are:")
    w.BUL([
        "**Asynchronous perception and capture-time registration (stage A).** The SLAM computer streams, with every pose, "
        "the reference keyframe used to track it; the recognition computer registers each late 6D result with the SLAM pose "
        "interpolated at its capture time and expresses it relative to the reference keyframe of that time, waiting when the "
        "pose has not arrived. Capture timestamps, clocks, and the camera–camera extrinsic are aligned jointly.",
        "**Keyframe-anchored spatial object mapping (stage B).** Each object is stored as 𝒪_{i} = {ID_{i}, K_{i}, "
        "T_{K_{i}←O_{i}}, p_{i}, C_{i}}. A loop closure updates only keyframe poses, re-sent as keyframe lists, so all objects "
        "follow the correction while T_{K_{i}←O_{i}} stays unchanged; existence probabilities keep objects through absences.",
        "**A verifiable zero-shot recognizer** that feeds stage A with few wrong answers at about 2 Hz, and an **evaluation** "
        "on YCB-Video [8], YCB-M [9], and indoor runs that separates the effect of capture time, keyframe-relative storage, "
        "and the choice of keyframe.",
    ])

    # ================================================================ 2 Related work
    w.H1(2, 'Related Work')
    w.P("**Zero-shot 6D pose estimation.** CNOS [10] and SAM-6D [3] segment the image with SAM [11] or FastSAM [12], "
        "describe each proposal with DINOv2 [13] features, and match it against templates rendered from the CAD model. "
        "FoundPose [4], GigaPose [5], and FoundationPose [14] follow related designs. These methods output camera-frame "
        "poses for single images. Open-vocabulary detectors such as YOLO-World [15] give class-level boxes from text "
        "prompts in tens of milliseconds; we use one as the proposal generator.")
    w.P("**Object SLAM and object maps.** SLAM++ [16], CubeSLAM [17], and QuadricSLAM [18] use objects as landmarks to "
        "improve localization; CosyPose [19] aggregates multi-view object poses; Fusion++ [6], POCD [20], and Khronos [21] "
        "maintain object maps with existence or change models. We do not feed objects back into SLAM; SLAM estimation is "
        "unchanged, and CAD-defined objects follow its corrections.")
    w.P("**Asynchronous perception.** MaskFusion [22] runs segmentation asynchronously on a separate GPU of one computer "
        "with one camera and one clock. Edge-assisted mobile AR [23] offloads detection and compensates the returned boxes "
        "by image tracking, without a map. Transform libraries [7] interpolate poses at a requested time but do not revisit "
        "stored results after the trajectory is optimized. To the best of our knowledge, no prior system binds late, "
        "off-device recognition results to the SLAM keyframe of their capture time.")

    # ================================================================ 3 Overview
    w.H1(3, 'Overview')
    w.span(w.FIG, 'fig_v3_times.png', 1,
           "Overview of the method. Stage A (asynchronous perception and capture-time registration): 1, sensor input on the "
           "cart; SLAM stream of poses and keyframes; 2, zero-shot 6D estimation with text-prompted proposals and "
           "verification gates; 3, asynchronous timing; 4, capture-time registration to the reference keyframe, producing an "
           "anchored object observation. Stage B (keyframe-anchored spatial object mapping): 5, object state map; 6, "
           "loop-closure-aware update. 7, result. All images, tables, and values are from recorded runs: panels 1–5 and 7 "
           "from the office figure-eight run, panel 6 from the large figure-eight run with a 3.05 m loop correction.", FULL_W)
    w.P("Figure 1 shows the method. Two computers are connected by a wired LAN. The SLAM computer runs ORB-SLAM3 (or "
        "HDL-Graph-SLAM) on its own RGB-D camera (or LiDAR); the recognition computer has a GPU and a second RGB-D camera. "
        "Stage A turns each 6D result into an *anchored object observation*: the object label, its capture time t_{c}, the "
        "reference keyframe K_{ref}(t_{c}), the relative pose T_{K←O}, a confidence, and its mask. Stage B accumulates these "
        "observations into a keyframe-anchored object state map and keeps the map consistent when SLAM corrects its "
        "keyframes. Notation: T_{A←B} maps coordinates from frame B to frame A; W is the map (world) frame; S and R are the "
        "SLAM and recognition cameras; X = T_{S←R} is the camera–camera extrinsic.")

    # ================================================================ 4 Stage A
    w.H1(4, 'Asynchronous Perception and Capture-Time Registration')
    w.H2('4.1', 'Sensor platform and time base')
    w.FIG('fig_hardware.jpg', 2, "Data-collection cart with two hardware-synchronized RGB-D cameras (SLAM and "
          "recognition) and a 16-channel LiDAR.", Cm(5.2))
    w.P("The cart (Fig. 1, panel 1; Fig. 2) carries two RGB-D cameras (640 × 480, 30 Hz), hardware-synchronized by a sync cable, and "
        "a 16-channel LiDAR (10 Hz). The capture time of every frame is the timestamp the camera attaches to the frame "
        "metadata; recorder association stamps deviated from the true frame times by −2.1 s to +3.8 s because of frame "
        "drops, whereas metadata timestamps agree within 1.4 ms. The clock offset between the computers is measured by UDP "
        "round trips,")
    w.EQ("Δ = t_{r} − ( t_{0} + (t_{1} − t_{0}) / 2 )", 1)
    w.P("with an error bound of half the round-trip time (3.08 ms measured). The remaining offset τ and the extrinsic X are "
        "estimated together from a calibration drive in three stages chosen by observability under planar motion: gravity "
        "from each trajectory plane; planar translation, yaw, and τ from relative motions over 0.5–4 s windows (AX = XB, "
        "Huber-weighted, τ scanned); and the height difference from the floor plane in each depth image. On the office "
        "figure-eight run the fit uses 7596 pairs with a median residual of 1.05 cm and τ = +2 ms.")

    w.H2('4.2', 'SLAM as a pose and keyframe stream')
    w.P("SLAM is an existing component; we change neither its estimation nor its map. Two additions make it usable by "
        "another computer. First, tracking stays within the camera period: with 2000 ORB features, result-preserving "
        "changes (removing copies for a disabled viewer, in-place aggregation, POPCNT descriptor distance, fewer locks, "
        "parallel pyramid extraction) reduce the mean tracking time from 24.7 ms to 15.4 ms and over-period frames from 170 "
        "to 0 of 7142, with the loop-return error unchanged (26.6 cm versus 26.5 cm, Fig. 3). Second, for every frame the "
        "SLAM computer sends a pose message (t, T_{W←S}, K_{ref}, T_{W←K_{ref}}), read in the tracking thread so that pose "
        "and reference keyframe refer to the same instant, and it sends the list of all live keyframes every 1 s, right "
        "after each map change, and again 0.5 s and 3 s later, because global bundle adjustment finishes seconds after the "
        "loop closure. A keyframe missing from a list is treated as culled.")
    w.FIG('fig_slam_timing.png', 3, "ORB-SLAM3 with 2000 features before and after the result-preserving optimizations.")

    w.H2('4.3', 'Zero-shot 6D estimation with verification')
    w.P("The recognizer (Fig. 1, panel 2; Fig. 4) modifies SAM-6D [3] so that it can run at about 2 Hz and refuse to answer. "
        "All object prompts (e.g., \"milk carton\", \"white jug\") are given to YOLO-World [15] in one pass, and MobileSAM "
        "[24] segments each box. Each candidate must then pass four gates in order: (1) semantic, the mean cosine of the five "
        "best template DINOv2 class tokens ≥ 0.35; (2) appearance, the mean best-match cosine of mid-layer patch features "
        "inside the mask ≥ 0.605; (3) color, the Bhattacharyya similarity of hue–saturation histograms ≥ 0.12; and (4) "
        "exclusive assignment of each box to one object. The SAM-6D pose estimation model then forms 6000 → 300 hypotheses, "
        "and a pose is accepted only if it passes silhouette overlap (IoU ≥ 0.42), feature agreement (≥ 0.45), and "
        "consensus (enough hypotheses within 20° and 25 mm), and again after refinement.")
    w.P("In the example frame of Fig. 1, YOLO-World returns 7 boxes (13.5 ms); four objects pass all gates and three "
        "candidates are rejected by the semantic gate (scores 0.26–0.32). For the milk carton, 267 of 300 hypotheses pass "
        "the silhouette check, 250 the feature check, and 153 form the consensus cluster; the accepted pose has IoU 0.87. "
        "On this laptop GPU the proposal and gate stage takes 84 ms and pose estimation for four objects 1.61 s.")
    w.span(w.FIG, 'fig_pipeline.png', 4,
           "Recognizer compared with SAM-6D. (A) Text-prompted proposals and sequential gates replace segment-everything "
           "proposals and weighted-sum scoring. (B) Pose verification is inserted before and after refinement.", FULL_W)

    w.H2('4.4', 'Asynchronous timing')
    w.P("The two streams run at different rates (Fig. 1, panel 3). SLAM produces a pose for every frame at 30 Hz; the "
        "recognizer processed 265 images in 128 s (2.1 Hz) on the office run. Each 6D result keeps its capture time t_{c} "
        "and arrives at t_{a} = t_{c} + Δt, with Δt between 0.7 s and 6.6 s across runs (median 1.57 s on the office run, "
        "1.45 s over 61 live runs). Pose messages, keyframe lists, and results travel over the LAN and can arrive out of "
        "order; in particular, the SLAM pose for t_{c} may arrive after the result.")

    w.H2('4.5', 'Capture-time registration')
    w.P("Registration (Fig. 1, panel 4) has three steps. (1) The SLAM pose at t_{c} is interpolated between the two "
        "bracketing poses (SLERP for rotation, linear for translation); results whose bracketing poses are more than 0.25 s "
        "apart are discarded, and results whose poses have not yet arrived wait in a queue for up to 20 s. (2) The reference "
        "keyframe K = K_{ref}(t_{c}) is the one sent with the bracketing pose nearest to t_{c}, together with its pose "
        "T_{W←K}(t_{c}) from the same message. (3) The extrinsic is applied and the object is expressed in that keyframe:")
    w.EQ("T_{K←O} = T_{W←K}(t_{c})^{−1} · T_{W←S}(t_{c}) · X · T_{R←O}", 2)
    w.P("Because T_{W←S}(t_{c}) and T_{W←K}(t_{c}) come from the same tracking state, (2) is unaffected by any correction "
        "applied to K after t_{c}, including corrections made while the result was in flight. Anchoring to the keyframe "
        "current at *arrival* would pair a pre-correction camera pose with a post-correction keyframe pose. The output is "
        "the anchored object observation of Fig. 1 (label, K, T_{K←O}, t_{c}, confidence, mask).")

    # ================================================================ 5 Stage B
    w.H1(5, 'Keyframe-Anchored Spatial Object Mapping')
    w.H2('5.1', 'Keyframe-anchored object state map')
    w.P("Each object in the map (Fig. 1, panel 5) is a record")
    w.EQ("𝒪_{i} = { ID_{i}, K_{i}, T_{K_{i}←O_{i}}, p_{i}, C_{i} }", 3)
    w.P("with identifier ID_{i}, anchor keyframe K_{i}, relative pose T_{K_{i}←O_{i}}, existence probability p_{i}, and "
        "class C_{i}, plus its observation history. Its pose in the map frame is never stored; it is computed when needed as")
    w.EQ("T_{W←O_{i}} = T_{W←K_{i}} · T_{K_{i}←O_{i}}", 4)
    w.P("A new observation is associated with an object of the same class within 0.15 m, nearest first and one-to-one per "
        "frame. Rotation is excluded from association because zero-shot estimators often return poses flipped by 180° for "
        "box-like objects. Position is the running mean of observations and orientation the mode of observations within "
        "30°. The existence probability drives the states tentative (p = 0.7 at creation, hidden), active (≥ 2 "
        "observations and p > 0.6), lost (p < 0.3), remembered (p < 0.05 with ≥ 5 observations), and deleted. A miss lowers "
        "p by Bayes' rule only if the object projects into the view *and* lies within the recognizer's detection range; "
        "lost and remembered objects stay association candidates, so a returning object regains its identifier. Of two "
        "same-class objects, one with five times fewer observations is hidden as a duplicate. On the office run, the map "
        "ends with eight active objects (14–29 observations each) and two hidden tentative ones.")

    w.H2('5.2', 'Loop-closure-aware update')
    w.P("When SLAM corrects its keyframes (Fig. 1, panel 6), the next keyframe list carries the corrected poses "
        "T′_{W←K_{i}}, and every object is updated by")
    w.EQ("T′_{W←O_{i}} = T′_{W←K_{i}} · T_{K_{i}←O_{i}}", 5)
    w.P("Only keyframe poses change; T_{K_{i}←O_{i}} is never modified, so all objects anchored to a keyframe move with it. "
        "An object whose keyframe was culled is re-anchored to the nearest live keyframe of the previous list, and the map can "
        "be rebuilt from the observation history whenever keyframe poses change (75 ms for 308 frames). For display, all map "
        "objects are projected at 15 Hz into a buffered image with the SLAM pose of that image's time, so drawing is "
        "independent of the recognition rate.")

    # ================================================================ 6 Setup
    w.H1(6, 'Experimental Setup')
    w.P("**Runs.** Recorded sessions are replayed in real time (1×) on the two computers, each replaying its own sensor "
        "stream. The SLAM computer is a laptop with an Apple M4 Pro processor; the recognition computer has an Intel Core "
        "i9-12900HK and an NVIDIA GeForce RTX 3080 Ti Laptop GPU [verify per run]. SLAM always builds its map from scratch. "
        "Indoor runs with eight objects: the dark figure-eight run (284 s), the 91 s run (96 map corrections), the bright-hall "
        "figure-eight run (151 s), the large figure-eight run (315 s; one repeat with a 3.05 m loop correction is used in "
        "Fig. 1), and the office figure-eight run (128 s). Public data: YCB-Video [8] (12 test videos) and YCB-M [9] (32 "
        "scenes). Single-image comparisons use an NVIDIA GeForce RTX 5090 [verify].")
    w.TAB(1, "Data used in the experiments.",
          ['Run / dataset', 'Length', 'Used for'],
          [['Dark figure-eight', '284 s', 'loop closure, 18 runs'],
           ['91 s run', '91 s, 96 corr.', 'anchoring, re-association'],
           ['Bright-hall figure-eight', '151 s', 're-association'],
           ['Large figure-eight', '315 s, 40 runs', 'loop update (3.05 m)'],
           ['Office figure-eight', '128 s', 'Fig. 1, timing, map'],
           ['YCB-Video [8]', '12 videos', 'recognition, display'],
           ['YCB-M [9]', '32 scenes', 'recognition, display']],
          [S(3.1), S(2.3), S(2.8)])
    w.P("**Metrics.** *Deviation* is the distance of an observation, registered in the final map, from the median position "
        "of its object. A pose is correct if ADD-S is below 10% of the object diameter; BOP AR uses the official toolkit; "
        "the *on-screen match rate* is the share of visible ground-truth objects whose displayed box is correct.")

    # ================================================================ 7 Results
    w.H1(7, 'Results')
    w.H2('7.1', 'Capture-time registration and keyframe anchoring')
    w.P("We replay 61 recorded live ORB-SLAM3 runs (8964 results; capture-to-arrival latency median 1.45 s, maximum "
        "6.6 s) with poses, keyframe updates, and results in their actual arrival order, and register the same results in "
        "four ways: (a) camera pose at arrival; (b) capture-time pose stored in the map frame; (c) capture-time pose anchored "
        "to the keyframe current at arrival; (d) proposed, anchored to the capture-time reference keyframe (Table 2, Fig. 5).")
    w.span(w.TAB, 2, "Deviation of repeated observations in the final map (61 runs).",
           ['Registration', 'Median (cm)', '90th pct. (cm)', '> 10 cm (%)', '90th pct., corrected (cm)', '> 10 cm, corrected (%)'],
           [['(a) Arrival-time pose', '9.67', '31.94', '48.2', '294.0', '63.8'],
            ['(b) Capture time, map frame', '3.08', '9.70', '9.2', '299.8', '30.8'],
            ['(c) Capture time, arrival-time keyframe', '1.59', '5.36', '3.6', '7.80', '5.4'],
            ['(d) Capture-time keyframe (proposed)', '1.51', '5.21', '3.5', '5.81', '1.5']],
           [S(5.6), S(1.9), S(2.2), S(1.9), S(2.8), S(2.6)],
           notes="Corrected: the 130 results whose wait spanned a map correction larger than 2 cm.",
           bold_cells={(4, j) for j in range(1, 6)})
    w.span(w.FIG, 'fig_registration_ablation.png', 5, "Registration strategies (a)–(d) on 61 replayed live runs.", FULL_W)
    w.P("Capture time removes most of the error (a → b). Without corrections in flight, (c) and (d) are similar, but among "
        "the 130 results whose wait spanned a map correction only (d) stays consistent (30.8% → 5.4% → 1.5% above 10 cm). "
        "The office run shows the same pattern with small corrections (maximum 11 cm): over 182 results, the share above "
        "10 cm is 63.7%, 13.7%, 7.7%, and 6.6% for (a)–(d), and among the 7 results with a correction in flight it is 57% "
        "for (b) and (c) but 14% for (d). Across a loop closure in the dark figure-eight run, the position difference of "
        "the same object before and after the loop falls from 5.4 cm to 0.8 cm. The waiting queue raises the share of "
        "results placed on the map from 46% to 100%.")

    w.H2('7.2', 'Loop-closure update and persistent memory')
    w.P("In the large figure-eight run of Fig. 1 (panel 6), a loop closure moved keyframes by up to 3.05 m; storing the "
        "same 119 results in the map frame leaves 68 more than 10 cm from the final object positions, versus 6 with "
        "keyframe anchoring. In the 91 s run, 96 corrections of up to 35 cm occur; without anchoring, three objects leave "
        "copies 30–32 cm away (11 locations for 8 objects). Persistent memory re-associates 72 of 74 reappearing objects in "
        "two indoor runs, versus 34 when association is done in camera coordinates (Table 3); the two failures occur before "
        "a loop closure, when drift displaces the object beyond the association radius.")
    w.TAB(3, "Re-association in indoor runs (object map / without map).",
          ['', 'Bright hall, 151 s', '91 s run'],
          [['Reappearances (out of view)', '35 (18)', '39 (21)'],
           ['Re-associated', '35 / 22', '37 / 12'],
           ['Identifiers created (8 objects)', '8 / 51', '10 / 68'],
           ['Same-class duplicates after run', '0 / 17', '0 / 19']],
          [S(3.6), S(2.4), S(2.2)])
    w.P("Majority-vote orientation reduces displays off by more than 20° or 50 mm from 23.9% to 9.3%, duplicate suppression "
        "hides high-confidence misreadings (observation ratios 65:3, 72:3, 68:4), and range-aware miss counting raises the "
        "display ratio of a small mug and can from 37%/46% to 83%/82%.")

    w.H2('7.3', 'Recognition')
    w.P("Table 4 compares SAM-6D with our recognizer using text prompts, and with ground-truth boxes substituted for the "
        "text detector. Our recognizer is correct in 94%–100% of its answers, produces at most one seventh of the wrong "
        "answers per image, and is 3.7–5.6 times faster; on YCB the text detector misses objects, so recall and BOP AR are "
        "lower than those of SAM-6D (78.7 versus 57.8 on YCB-V), which the method compensates by persistent memory rather "
        "than by per-frame recall. Adding the gates one at a time on 959 human-labeled cells raises F1 from 0.692 to 0.809 "
        "and reduces false positives from 310 to 28 (Fig. 6).")
    w.span(w.TAB, 4, "Single-image comparison without SLAM (one answer per object per image).",
           ['Dataset / method', 'Found (%)', 'Correct answers (%)', 'Wrong per image', 'BOP AR', 'Time (ms)'],
           [['YCB-V · SAM-6D', '90.2', '90.9', '0.41', '78.7', '1613'],
            ['YCB-V · ours, text', '63.2', '98.6', '0.04', '57.8', '433'],
            ['YCB-V · ours, GT box', '80.9', '100.0', '0.00', '74.5', '528'],
            ['YCB-M · SAM-6D', '79.1', '80.0', '1.01', '52.3', '2197'],
            ['YCB-M · ours, text', '42.6', '94.0', '0.14', '29.7', '395'],
            ['Indoor · SAM-6D', '52.4', '43.0', '3.39', '33.3', '2884'],
            ['Indoor · ours, text', '51.8', '94.7', '0.14', '37.9', '557']],
           [S(4.4), S(2.0), S(2.8), S(2.4), S(2.0), S(2.0)],
           notes="GT, ground truth; AR, average recall. YCB-V 900, YCB-M 925, indoor 917 images (pseudo ground truth).")
    w.FIG('fig_gate_ablation.png', 6, "Adding gates one at a time on 959 human-labeled cells: weighted sum (Sum), color "
          "(+C), appearance (+A), exclusive assignment (+E).")

    w.H2('7.4', 'Real-time behavior and display')
    w.P("On the office run, SLAM tracks all 3844 input frames (11.6 ms each, 0 dropped) while recognition runs at 2.1 Hz "
        "with results arriving 1.57 s late (median), and no result remains unplaced. Drawing the object map at the "
        "display-time pose raises the on-screen match rate from 43.3% to 85.7% on YCB-Video and from 3.2% to 72.7% on "
        "YCB-M, where SAM-6D processes only 8% of frames (Fig. 7).")
    w.P("Table 5 shows the final object maps of live YCB-Video playback with the same SLAM and map: SAM-6D results "
        "produce eight duplicates over 12 scenes and a maximum position error of 85.1 mm, whereas the verified recognizer "
        "produces none, and wrongly drawn boxes fall from 0.266 to 0.001 per frame.")
    w.TAB(5, "Final object maps, YCB-Video live playback (12 scenes, 55 objects).",
          ['', 'SAM-6D', 'Ours (GT box)', 'Ours (text)'],
          [['Registered / GT', '55/55', '51/55', '42/55'],
           ['Duplicates', '8', '0', '0'],
           ['Pos. error med. (mm)', '7.1', '5.2', '5.4'],
           ['Pos. error max (mm)', '85.1', '23.8', '76.6'],
           ['Wrong boxes / frame', '0.266', '0.001', '0.001']],
          [S(2.9), S(1.7), S(1.9), S(1.7)])
    w.span(w.FIG, 'fig_display_match.png', 7, "On-screen match rate: last result held in camera coordinates versus the "
           "object map drawn at the display-time SLAM pose.", FULL_W)

    w.H2('7.5', 'Limitations')
    w.P("Anchoring follows SLAM corrections, including wrong ones: in 40 runs of the large figure-eight with identical "
        "input, two loop closures damaged an accurate map. The recognizer is weak under occlusion (3% versus 46% found for "
        "objects 30%–50% visible on YCB-Video), and acceptance near the thresholds varies with the random hypotheses of "
        "pose estimation (in the example frame, the dinosaur doll passes the final feature check in 2 of 10 seeds). Scenes "
        "with several objects of one class were not evaluated, and indoor runs lack ground-truth object poses.")

    w.H1(8, 'Discussion')
    w.P("The two stages answer three questions that are often conflated. *Which time* registers the result: replacing "
        "arrival time with capture time accounts for most of the improvement (9.67 → 3.08 cm), because the platform moves "
        "while the result is computed. *Which frame* stores the object: storing it relative to a keyframe instead of the "
        "map frame is what lets it follow later corrections. *Which keyframe*: the keyframe current at arrival is adequate "
        "most of the time but fails when a correction happens in flight, because it pairs a pre-correction camera pose with a "
        "post-correction keyframe pose; the capture-time keyframe avoids this by construction (30.8% → 5.4% → 1.5% above "
        "10 cm). Since zero-shot recognizers are slow and their latency varies, corrections in flight are not rare: 130 of "
        "8964 results in our runs, concentrated around loop closures, where an object map is most likely to break.")
    w.P("The SLAM side only exposes information it already maintains, the reference keyframe of each pose and the poses of "
        "all keyframes, and schedules their transmission. The method therefore applies to other keyframe- or graph-based "
        "SLAM systems, as shown with ORB-SLAM3 and HDL-Graph-SLAM, and to other recognizers that output camera-frame poses. "
        "The recognizer was designed for this setting: in an object map a wrong answer leaves a persistent error, whereas a "
        "missed detection is recovered when the object is seen again, so precision is favored over per-frame recall.")
    w.H1(9, 'Conclusion')
    w.P("This study maps late zero-shot 6D pose estimates consistently on an asynchronous SLAM map in two stages: "
        "capture-time registration, which pairs each result with the SLAM pose and reference keyframe of its capture time, "
        "and a keyframe-anchored object state map, in which a loop closure updates only keyframe poses while each object's "
        "relative pose T_{K←O} stays fixed. The separation of capture time, keyframe-relative storage, and keyframe choice "
        "shows that the last matters exactly when a correction happens in flight. Future work includes detecting wrong loop "
        "closures, occlusion-aware verification, and ground-truth object poses for indoor runs.")
    refs(d)
    add_page_numbers(d)
    d.save(path)


if __name__ == '__main__':
    out = os.path.join(B.ROOT, 'etrij'); os.makedirs(out, exist_ok=True)
    build(os.path.join(out, 'ETRIJ_manuscript_v0.2.docx'))
    print('abstract chars:', len(ABSTRACT))
