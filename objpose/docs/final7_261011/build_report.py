#!/usr/bin/env python3
"""build_report.py — 2026-10-11 확정 구성 일곱 가지(자세 추정 다섯 + FM-Fusion + DualMap)의 실시간 객체 지도 결과를 한 보고서로 정리한다.
  final7_report.html / _local.html / _notion.html, --pptx 를 주면 ~/콜로키움/전체_실험결과.pptx 뒤에 요약 쪽을 덧붙인다(제목 접두 "종합").
  입력  ObChange SmallRoom  objpose/output/rt_obchange_261011/{scores.json(score_designs.py), scores_base.json(score_base_same.py), designs/, c<k>/<방문>/run.json}
        수집 영상 8개       objpose/output/live_rt_261011/designs/ (지도 설계를 실행 중에 계산한 결과), FM-Fusion · DualMap 은 objpose/output/lab_combo_261010/objlist.json
        그림                media/pose_*.jpg (objpose/tools/analysis/render_pose_map.py --config 8 --variant V0, OBJPOSE_OB_OUT=rt_obchange_261011 OBJPOSE_LAB_OUT=live_rt_261011)
    python3 objpose/docs/final7_261011/build_report.py [--pptx]"""
import argparse, importlib.util, json, re, sys
from pathlib import Path
import numpy as np
HERE = Path(__file__).parent; REPO = Path(__file__).resolve().parents[3]; MED = HERE / "media"; OB = REPO / "objpose/output/rt_obchange_261011"; LAB = REPO / "objpose/output/live_rt_261011"


def mod(name, path):
    spec = importlib.util.spec_from_file_location(name, path); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m


B = mod("objmap_rep", REPO / "objpose/docs/objmap_eval_261009/build_report.py"); Doc, REF = B.Doc, B.REF; J = lambda p: json.loads(Path(p).read_text())
CS = [2, 9, 8, 7, 6]; NO = {2: "①", 9: "②", 8: "③", 7: "④", 6: "⑤"}
NAME = {2: "① 원본 (FastSAM-x + 총점)", 9: "② 관문 비교 (FastSAM-x + 관문)", 8: "③ 제안 방법 (YOLO-World + MobileSAM + 관문)", 7: "④ 완화 비교 (YOLO-World + MobileSAM + 완화 관문)", 6: "⑤ 총점 효과 (YOLO-World + MobileSAM + 총점)"}
SPEC = {2: ("화면 전체 분할 (FastSAM-x)", "FastSAM-x", "원본 SAM-6D 총점, 물체마다 1위, 0.5 이상"), 9: ("화면 전체 분할 (FastSAM-x) 의 상자", "상자 안을 MobileSAM 으로", "제안 관문"),
        8: ("문구 상자 (YOLO-World)", "MobileSAM", "제안 관문"), 7: ("문구 상자 (YOLO-World)", "MobileSAM", "제안 관문 완화 (의미 기준 0.25, 2 · 3위 물체에게도 상자 배정)"), 6: ("문구 상자 (YOLO-World)", "MobileSAM", "원본 SAM-6D 총점, 물체마다 1위, 0.5 이상")}
VK = ["V0", "1", "1n", "2", "3", "1+2", "1n+2", "1+3", "1n+3", "2+3", "1+2+3", "1n+2+3"]
VN = {"V0": "지금 방식", "1": "설계 1", "1n": "설계 1′", "2": "설계 2", "3": "설계 3", "1+2": "1+2", "1n+2": "1′+2", "1+3": "1+3", "1n+3": "1′+3", "2+3": "2+3", "1+2+3": "1+2+3", "1n+2+3": "1′+2+3"}
f1 = lambda x: f"{x:.1f}"
# 사진 한 장 평가(objpose/docs/ycb_combo_261008 8절, 전체 이미지, 물체를 모르는 조건 = 21종 질의, 자세 검증 켬): (찾은 비율 %, 정답 비율 %, 오답 / 장). ② 는 전체 이미지 실행이 없다
STILL = {"YCB-Video 900장": {2: (76.3, 87.3, 0.510), 6: (75.7, 91.3, 0.330), 7: (60.4, 83.9, 0.531), 8: (46.1, 86.8, 0.320)}, "YCB-M 925장": {2: (62.3, 90.6, 0.330), 6: (61.6, 90.4, 0.332), 7: (55.5, 83.2, 0.569), 8: (43.5, 87.8, 0.307)}}


def agg(rows):
    gt = sum(r["gt"] for r in rows); tp = sum(r["tp"] for r in rows); nm = sum(r["named"] for r in rows); t5 = sum(r["tp50"] for r in rows); io = [x for r in rows for x in r["iou_tp"]]
    p = tp / max(nm, 1); rc = tp / max(gt, 1); return {"gt": gt, "tp": tp, "fp": nm - tp, "fn": gt - tp, "P": 100 * p, "R": 100 * rc, "F": 200 * p * rc / max(p + rc, 1e-9), "tp50": t5, "iou": float(np.mean(io)) if io else float("nan")}


def lab_stat(runs, v):
    k = e = 0; sp = []
    for x in runs:
        nm = [m["name"] for m in x[v]]; k += len(set(nm)); e += len(nm) - len(set(nm)); sp += [m["support"] for m in x[v] if m["support"] is not None]
    return {"kinds": k, "extra": e, "sup": float(np.mean(sp)) if sp else float("nan")}


def online(runs):
    g = lambda f: [f(x["online"]) for x in runs]; al = lambda k, s: [x["online"]["alone_ms"][k][s] for x in runs]
    return {"passes": sum(g(lambda o: o["passes"])), "all_med": float(np.median(g(lambda o: o["pass_ms"]["median"]))), "all_max": max(g(lambda o: o["pass_ms"]["max"])), "lag_med": float(np.median(g(lambda o: o["lag_ms"]["median"]))), "lag_max": max(g(lambda o: o["lag_ms"]["max"])),
            **{f"{k}_{s}": (float(np.median(al(k, s))) if s == "median" else max(al(k, s))) for k in ("V0", "1n", "1n+2", "1n+3", "1n+2+3") for s in ("median", "p95", "max")}}


def load():
    S = J(OB / "scores.json"); SB = J(OB / "scores_base.json"); T = {(c, v): agg([r for r in S if r["config"] == c and r["variant"] == v]) for c in CS for v in VK}; A = {v: agg([r for r in S if r["variant"] == v and r["config"] in CS]) for v in VK}
    BT = {m: agg([r for r in SB if r["method"] == m]) for m in ("FM-Fusion", "DualMap")}
    for m in BT: BT[m]["all"] = sum(r["all_objects"] for r in SB if r["method"] == m)
    OD = {c: [J(f) for f in sorted((OB / "designs").glob(f"c{c}_smallroom_s*.json"))] for c in CS}; LD = {c: [J(f) for f in sorted((LAB / "designs").glob(f"c{c}_26*.json"))] for c in CS}
    assert all(len(OD[c]) == 5 and len(LD[c]) == 8 for c in CS), {c: (len(OD[c]), len(LD[c])) for c in CS}
    LT = {(c, v): lab_stat(LD[c], v) for c in CS for v in VK}; LA = {v: lab_stat([x for c in CS for x in LD[c]], v) for v in VK}
    lo = J(REPO / "objpose/output/lab_combo_261010/objlist.json"); LB = {}
    for m in ("FM-Fusion", "DualMap"):
        named = [[o for o in lo[s][m] if o.get("name")] for s in lo]; LB[m] = {"kinds": sum(len({o["name"] for o in n}) for n in named), "extra": sum(sum(bool(o.get("extra")) for o in n) for n in named), "all": sum(len(lo[s][m]) for s in lo)}
    RUN = {c: [J(f) for f in sorted(OB.glob(f"c{c}/smallroom_s*/run.json"))] for c in CS}
    return dict(T=T, A=A, BT=BT, OD=OD, LD=LD, LT=LT, LA=LA, LB=LB, RUN=RUN, ON_OB={c: online(OD[c]) for c in CS}, ON_LAB={c: online(LD[c]) for c in CS})

def repeats(T):
    """같은 인식 결과(③)에 지도 생성 방식만 바꾼 3회 반복(objpose/output/mapgen_261012) + ③ 대 ⑤ SmallRoom 반복."""
    B = REPO / "objpose/output/mapgen_261012"; ob = json.loads((B / "agg_obchange.json").read_text()); yc = json.loads((B / "agg_ycbv.json").read_text())
    rg = lambda xs, f="{:.0f}": (f.format(min(xs)) if min(xs) == max(xs) else f.format(min(xs)) + "~" + f.format(max(xs))); mean = lambda xs: sum(xs) / len(xs)
    L = [("M0 단순 누적", "M0 단순 누적", "M0"), ("M1 기존 갱신 (위치 평균, 방향 다수결)", "M1 제안", "M1"), ("M2 합의로 포즈 결정", "M2 합의(1n)", "M2"), ("M3 깊이로 다듬기", "M3 깊이 다듬기(2)", "M3"), ("M4 FM-Fusion 의 지도 생성", "M4 FM-Fusion", "M4"), ("M5 DualMap 의 지도 생성", "M5 DualMap", "M5")]
    OB = [[n, rg([x["tp"] for x in ob["all|" + k]]), rg([x["fp"] for x in ob["all|" + k]]), f"{mean([x['f1'] for x in ob['all|' + k]]):.1f} ({rg([x['f1'] for x in ob['all|' + k]], '{:.1f}')})", rg([x["tp50"] for x in ob["all|" + k]]), f"{mean([x['iou'] for x in ob['all|' + k]]):.3f}"] for n, k, _ in L]
    YC = [[n, rg([x["tp"] for x in yc[m]]), rg([x["fp"] for x in yc[m]]), rg([x["f1"] for x in yc[m]], "{:.1f}"), f"{mean([x['iou'] for x in yc[m]]):.3f}", rg([x["adds_med"] for x in yc[m]], "{:.1f}") if yc[m][0]["adds_med"] is not None else "자세 없음", rg([x["trans_med"] for x in yc[m]], "{:.1f}") if yc[m][0]["trans_med"] is not None else "자세 없음", rg([x["trans_p90"] for x in yc[m]], "{:.1f}") if yc[m][0]["trans_p90"] is not None else "자세 없음"] for n, _, m in L]
    def sm(f):
        c = {"tp": 0, "named": 0, "gt": 0}
        for r in json.loads(Path(f).read_text()):
            if r["variant"] == "V0" and r["scene"].startswith("smallroom"):
                for k in c: c[k] += r[k]
        p = c["tp"] / max(c["named"], 1); rc = c["tp"] / c["gt"]; return c["tp"], c["named"] - c["tp"], 200 * p * rc / (p + rc)
    g3 = [(T[(8, "V0")]["tp"], T[(8, "V0")]["fp"], T[(8, "V0")]["F"])] + [(x["tp"], x["fp"], x["f1"]) for x in ob["smallroom|M1 제안"]]
    g5 = [(T[(6, "V0")]["tp"], T[(6, "V0")]["fp"], T[(6, "V0")]["F"])] + [sm(B / f"gate/rep{k}/scores.json") for k in (2, 3)]
    GT = [["③ 제안 관문", len(g3), rg([x[0] for x in g3]), rg([x[1] for x in g3]), rg([x[2] for x in g3], "{:.1f}")], ["⑤ 총점 판정", len(g5), rg([x[0] for x in g5]), rg([x[1] for x in g5]), rg([x[2] for x in g5], "{:.1f}")]]
    N = json.loads((B / "agg_night.json").read_text()); fm = lambda xs: f"{mean(xs):.1f} ({rg(xs, '{:.1f}')})"; RM = ["smallroom", "office", "livingarea", "kitchen"]
    CN = {"① 원본": "① 원본 (화면 전체 분할 + 총점)", "② 전체 분할+관문": "② 화면 전체 분할 + 관문", "③ 제안": "③ 제안 (문구 상자 + 관문)", "⑤ 총점 판정": "⑤ 문구 상자 + 총점"}
    NB = [[CN[k], len(x), rg([v["all"]["tp"] for v in x]), rg([v["all"]["fp"] for v in x]), fm([v["all"]["f1"] for v in x])] + [f"{mean([v[r]['f1'] for v in x]):.1f}" for r in RM] for k, x in N["obchange"].items()]
    NY = [[CN[k.split("|")[0]].split(" (")[0] + (", 기존 갱신" if k.endswith("M1") else ", 합의"), rg([v["tp"] for v in x]), rg([v["fp"] for v in x]), fm([v["f1"] for v in x]), rg([v["adds_med"] for v in x], "{:.1f}"), rg([v["trans_med"] for v in x], "{:.1f}")] for k, x in N["ycbv"].items()]
    po = lambda x, m: rg([v[m] for v in x], "{:.1f}") if x[0][m] is not None else "자세 없음"
    YM = [[n, rg([v["tp"] for v in N["ycbm"][m]]), rg([v["fp"] for v in N["ycbm"][m]]), fm([v["f1"] for v in N["ycbm"][m]]), po(N["ycbm"][m], "adds_med"), po(N["ycbm"][m], "trans_med"), po(N["ycbm"][m], "trans_p90")] for n, _, m in L]
    return {"NB": NB, "NY": NY, "YM": YM, "N": N, "ymgt": N["ycbm"]["M1"][0]["gt"], "OB": OB, "YC": YC, "GT": GT, "ngt": ob["all|M1 제안"][0]["gt"], "ob": ob, "yc": yc, "g3": g3, "g5": g5, "rg": rg}


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--pptx", action="store_true"); a = ap.parse_args(); E = load(); T, A, BT, LT, LA, LB, RUN, ON_OB, ON_LAB = (E[k] for k in ("T", "A", "BT", "LT", "LA", "LB", "RUN", "ON_OB", "ON_LAB"))
    d = Doc(); ngt = T[(8, "V0")]["gt"]; P = T[(8, "V0")]; LP = LT[(8, "V0")]; RP = repeats(T); E["RP"] = RP
    bestF = max(CS, key=lambda c: T[(c, "V0")]["F"]); E["bestF"] = bestF
    d.h2("1. 결론")
    d.ul([f"<strong>{NAME[bestF]}</strong>과 지금 방식의 지도 생성을 쓴 조합이 두 데이터 모두에서 가장 좋았습니다. ObChange SmallRoom 에서 F1 {f1(T[(bestF, 'V0')]['F'])} %(맞음 {T[(bestF, 'V0')]['tp']}/{ngt}, 잘못 등록 {T[(bestF, 'V0')]['fp']}), 직접 수집한 영상 8개에서 {LT[(bestF, 'V0')]['kinds']}/64종 등록에 초과 등록 {LT[(bestF, 'V0')]['extra']}개입니다.",
          f"물체 판정을 원본 총점에서 제안 관문으로 바꾼 효과가 가장 큽니다. 같은 앞단에서 총점(⑤)은 잘못 등록 {T[(6, 'V0')]['fp']}개, 관문(③)은 {T[(8, 'V0')]['fp']}개이고 F1은 {f1(T[(6, 'V0')]['F'])} → {f1(T[(8, 'V0')]['F'])} % 입니다. 화면 전체 분할에서도 같은 방향입니다(① {f1(T[(2, 'V0')]['F'])} → ② {f1(T[(9, 'V0')]['F'])} %).",
          f"FM-Fusion 과 DualMap 은 맞은 수는 비슷하지만(각 {BT['FM-Fusion']['tp']}, {BT['DualMap']['tp']}/{ngt}) 이름이 붙은 잘못된 객체가 {BT['FM-Fusion']['fp']}, {BT['DualMap']['fp']}개여서 F1이 {f1(BT['FM-Fusion']['F'])}, {f1(BT['DualMap']['F'])} % 입니다.",
          f"지도 설계(합의로 합치기 등)는 이번 실시간 실행에서 지금 방식을 넘지 못했습니다. 설계 1′은 자세가 정확한 객체(겹침 0.5 이상)를 {A['V0']['tp50']} → {A['1n']['tp50']}개로 늘렸지만 맞은 수가 {A['V0']['tp']} → {A['1n']['tp']}개로 줄어 F1이 {f1(A['V0']['F'])} → {f1(A['1n']['F'])} % 로 내려갔습니다. 어젯밤 사후 계산에서 본 \"맞은 수는 그대로\"는 다시 나오지 않았습니다.",
          f"사진 한 장씩 평가한 10월 8일 실험에서는 순위가 반대였습니다. 총점 판정(⑤)이 제안 관문(③)보다 훨씬 많이 찾았고(YCB-Video {STILL['YCB-Video 900장'][6][0]} 대 {STILL['YCB-Video 900장'][8][0]} %) 오답은 같은 수준이었습니다. 제안 관문의 이득은 객체 지도 평가에서만 나타납니다(3절 끝).",
          f"반복 실행(6절)에서 관문의 효과는 유지됐습니다. SmallRoom 에서 제안 관문(③, {RP['GT'][0][1]}회)은 잘못 등록 {RP['GT'][0][3]}건, 총점 판정(⑤, {RP['GT'][1][1]}회)은 {RP['GT'][1][3]}건으로 범위가 겹치지 않고, 맞음은 {RP['GT'][0][2]} 대 {RP['GT'][1][2]} 로 같은 수준입니다.",
          f"같은 인식 결과(③)를 여러 지도 생성 방식에 넣어 3회 반복하면(6절), 지금 방식 계열(M1 · M2)이 FM-Fusion · DualMap 의 지도 생성보다 F1 이 높습니다. 합의로 포즈를 정하면(M2) 등록 수는 그대로이고 YCB-Video 위치 오차 중앙값이 {RP['YC'][1][6]} → {RP['YC'][2][6]} mm 로 줄었습니다. 다만 걸러내지 않는 단순 누적(M0)과 F1 은 차이가 없습니다.",
          "1~5절의 수치는 구성마다 한 번 돌린 결과이고, 6절이 3회 반복입니다."])
    d.h2("2. 무엇을 비교했나")
    d.table(["번호", "후보 생성", "마스크", "물체 판정", "지도 생성"], [[NAME[c].split(" (")[0], *SPEC[c], "지금 방식 + 설계 조합 11가지 (실행 중 계산)"] for c in CS] + [["⑥ FM-Fusion", "GroundingDINO 문구 상자", "SAM", "검출기의 이름", "자체 방식 (점 묶음)"], ["⑦ DualMap", "YOLO-World 문구 상자", "MobileSAM / FastSAM", "검출기의 이름 + CLIP", "자체 방식 (점 묶음)"]], cls="sum")
    d.ul(["①~⑤ 는 모두 특징 ViT-S, 자세 추정(PEM, 점 1024개), 자세 검증(윤곽 겹침 0.421 이상, 질감 0.450 이상, 합의 20° · 25 mm, 합의 비율 0.5 이상)을 켠 상태입니다. 다섯 구성의 실행 기록에서 검증이 켜진 것을 확인했습니다.",
          "② 는 FastSAM-x 가 낸 상자를 후보로 쓰고, 관문 계산에 쓰는 마스크는 그 상자 안에서 MobileSAM 으로 다시 만듭니다(제안 관문 코드가 상자를 입력으로 받기 때문).",
          "③ · ④ · ⑤ 는 후보 상자 규칙(점수 0.02 이상, 문구마다 상위 3개), 마스크, 특징이 같고 물체 판정만 다릅니다.",
          "데이터: 직접 수집한 영상 8개(노트북의 ORB-SLAM3 가 실시간으로 낸 카메라 위치, 정답 없음)와 ObChange SmallRoom 방문 5회(데이터셋 제공 카메라 위치, 10 fps 실시간 재생, 정답 물체 " + str(ngt) + "개).",
          "⑥ · ⑦ 은 10월 9 ~ 10일에 실시간 재생으로 돌린 결과를 그대로 썼습니다. 수집 영상의 카메라 위치는 이 PC 에서 돌린 ORB-SLAM3 결과여서 ①~⑤(노트북 실시간)와 실행이 다릅니다. 채점 규칙은 ①~⑤ 와 같게 다시 적용했습니다."])
    d.h2("3. ObChange SmallRoom (정답 있음)")
    d.p("정답 물체와 같은 이름의 객체가 관측 점 기준으로 0.25 이상 겹치면 맞음, 이름이 붙었는데 맞음이 아니면 잘못 등록입니다. ①~⑤ 는 지금 방식의 지도입니다.")
    row = lambda n, t: [n, t["tp"], t["fp"], t["fn"], t["P"], t["R"], t["F"], t["tp50"], f"{t['iou']:.3f}"]
    d.table(["구성", "맞음", "잘못 등록", "놓침", "정밀도 %", "재현율 %", "F1 %", "겹침 0.5 이상", "맞은 객체의 평균 겹침"], [row(NAME[c], T[(c, "V0")]) for c in CS] + [row("⑥ FM-Fusion", BT["FM-Fusion"]), row("⑦ DualMap", BT["DualMap"])], bars=(6,))
    d.ul([f"FM-Fusion 은 객체 {BT['FM-Fusion']['all']}개 가운데 {BT['FM-Fusion']['tp'] + BT['FM-Fusion']['fp']}개, DualMap 은 {BT['DualMap']['all']}개 가운데 {BT['DualMap']['tp'] + BT['DualMap']['fp']}개에 물체 목록의 이름이 붙었습니다. 이름이 목록에 없는 객체는 세지 않았습니다.",
          "처리한 프레임(방문 5회 합, 전체 " + str(sum(r["frames"] for r in RUN[8])) + "장): " + ", ".join(f"{NO[c]} {sum(r['processed'] for r in RUN[c])}장" for c in CS) + ". 화면 전체 분할을 쓰는 ① · ② 는 한 장에 시간이 더 걸려 처리한 장수가 적습니다."])
    d.h3("구성 사이의 비교")
    cmp = lambda a_, b_, what: [what, NO[a_] + " → " + NO[b_], f"{T[(a_, 'V0')]['tp']} → {T[(b_, 'V0')]['tp']}", f"{T[(a_, 'V0')]['fp']} → {T[(b_, 'V0')]['fp']}", f"{f1(T[(a_, 'V0')]['F'])} → {f1(T[(b_, 'V0')]['F'])}", f"{LT[(a_, 'V0')]['kinds']} → {LT[(b_, 'V0')]['kinds']}", f"{LT[(a_, 'V0')]['extra']} → {LT[(b_, 'V0')]['extra']}"]
    CMP = [cmp(2, 9, "총점 → 관문 (화면 전체 분할에서)"), cmp(6, 8, "총점 → 관문 (문구 상자에서)"), cmp(9, 8, "화면 전체 분할 → 문구 상자 (관문은 같음)"), cmp(8, 7, "관문 → 완화 관문"), cmp(2, 8, "원본 → 제안 방법")]; E["CMP"] = CMP
    d.table(["바꾼 것", "구성", "SmallRoom 맞음", "SmallRoom 잘못 등록", "SmallRoom F1 %", "수집 영상 등록 종류 (최대 64)", "수집 영상 초과 등록"], CMP)
    d.h3("사진 한 장 평가와 객체 지도 평가를 나란히")
    d.p("10월 8일 조합 실험은 같은 구성을 사진 한 장씩 평가했습니다(ADD-S 가 물체 지름의 10 % 미만이면 맞음, 물체 21종을 모두 묻는 조건, 자세 검증 켬). 그 결과와 이번 객체 지도 결과를 한 표에 놓았습니다.")
    SV, SM = STILL["YCB-Video 900장"], STILL["YCB-M 925장"]; g = lambda D, c, i, fmt: (fmt.format(D[c][i]) if c in D else "실행 없음")
    d.table(["구성", "한 장: YCB-Video 찾은 비율 %", "한 장: YCB-Video 오답 / 장", "한 장: YCB-M 찾은 비율 %", "한 장: YCB-M 오답 / 장", "지도: SmallRoom 맞음 (정답 " + str(ngt) + ")", "지도: SmallRoom 잘못 등록", "지도: SmallRoom F1 %", "지도: 수집 영상 초과 등록"],
            [[NAME[c], g(SV, c, 0, "{:.1f}"), g(SV, c, 2, "{:.3f}"), g(SM, c, 0, "{:.1f}"), g(SM, c, 2, "{:.3f}"), T[(c, "V0")]["tp"], T[(c, "V0")]["fp"], f1(T[(c, "V0")]["F"]), LT[(c, "V0")]["extra"]] for c in CS])
    d.ul([f"사진 한 장 평가에서는 총점 판정(① · ⑤)이 제안 관문(③)보다 약 30 %p 더 많이 찾았고(YCB-Video {SV[6][0]} 대 {SV[8][0]} %), 오답은 ⑤ 와 ③ 이 같은 수준이었습니다({SV[6][2]:.3f} 대 {SV[8][2]:.3f} / 장).",
          f"객체 지도 평가에서는 맞은 수가 같은 수준이고({T[(6, 'V0')]['tp']} 대 {T[(8, 'V0')]['tp']}), 잘못 등록은 총점 쪽이 세 배입니다({T[(6, 'V0')]['fp']} 대 {T[(8, 'V0')]['fp']}).",
          "두 평가는 세는 대상(프레임마다의 자세 대 영상이 끝난 뒤 지도에 남은 객체)과 데이터가 다릅니다. 순위가 뒤집힌 원인을 따로 확인한 실험은 없습니다.",
          "한 장 평가의 ③ 은 배포 설정의 원래 문구, ① · ⑤ 는 새 문구를 썼습니다. 판정과 문구가 함께 바뀐 비교입니다. ② 는 전체 이미지 실행이 없어 비웠습니다."])
    d.h2("4. 직접 수집한 영상 8개 (실제 환경, 정답 없음)")
    d.p("영상마다 물체 8종이 하나씩 놓여 있습니다. 등록된 종류가 많고(최대 64) 같은 종류의 초과 등록이 적을수록 좋습니다. 깊이 지지율은 지도의 자세대로 CAD 를 놓았을 때 깊이 영상의 표면과 맞는 점의 비율입니다.")
    d.table(["구성", "등록된 종류 (최대 64)", "초과 등록", "깊이 지지율"], [[NAME[c], LT[(c, "V0")]["kinds"], LT[(c, "V0")]["extra"], f"{LT[(c, 'V0')]['sup']:.2f}"] for c in CS] + [[f"⑥ FM-Fusion (객체 {LB['FM-Fusion']['all']}개)", LB["FM-Fusion"]["kinds"], LB["FM-Fusion"]["extra"], "해당 없음"], [f"⑦ DualMap (객체 {LB['DualMap']['all']}개)", LB["DualMap"]["kinds"], LB["DualMap"]["extra"], "해당 없음"]])
    d.h2("5. 객체 지도 생성 방식 비교 (실시간 실행 결과)")
    d.p("인식 결과가 한 프레임 도착할 때마다 그때까지의 자료만으로 열두 가지 지도를 모두 갱신했습니다. 설계는 인식에 영향을 주지 않으므로 한 번의 실행에서 함께 냅니다. 즉 아래 수치는 영상을 실시간으로 넣으면서 지도까지 그 자리에서 만든 결과이고, 실행이 끝난 뒤 기록으로 다시 계산한 것이 아닙니다.")
    d.h3("표 5-1. 구성 ①~⑤ 의 합 (지도 방식의 전반적인 경향)")
    d.table(["지도 생성 방식", "SmallRoom 맞음", "잘못 등록", "F1 %", "겹침 0.5 이상", "평균 겹침", "수집 영상 등록 종류 (최대 320)", "초과 등록", "깊이 지지율"], [[VN[v], A[v]["tp"], A[v]["fp"], A[v]["F"], A[v]["tp50"], f"{A[v]['iou']:.3f}", LA[v]["kinds"], LA[v]["extra"], f"{LA[v]['sup']:.2f}"] for v in VK], bars=(3,))
    d.h3("표 5-2. 제안 방법(③)만 (최종 조합을 고를 때 보는 표)")
    d.p("표 5-1 과 지표는 같고 대상이 다릅니다. 위는 다섯 구성을 더한 값, 아래는 제안 방법 하나의 값입니다.")
    d.table(["지도 생성 방식", "SmallRoom 맞음", "잘못 등록", "F1 %", "겹침 0.5 이상", "수집 영상 등록 종류 (최대 64)", "초과 등록"], [[VN[v], T[(8, v)]["tp"], T[(8, v)]["fp"], T[(8, v)]["F"], T[(8, v)]["tp50"], LT[(8, v)]["kinds"], LT[(8, v)]["extra"]] for v in VK], bars=(3,))
    d.ul(["설계 1(합의로 합치기): 객체마다 서로 가까운(25 mm, 20°) 자세의 가장 큰 무리만 써서 자세를 정합니다. 1′은 무리가 작은 객체를 빼는 규칙을 없앤 것입니다. 설계 2: 깊이 영상으로 확인하고 CAD 를 다시 맞춥니다. 설계 3: 검증에서 떨어진 후보도 증거로 쌓아 등록합니다.",
          f"합의는 자세가 맞을 때는 더 정확하게 만들지만(겹침 0.5 이상 {A['V0']['tp50']} → {A['1n']['tp50']}개), 무리를 잘못 고르면 평균보다 더 벗어나 맞음에서 빠집니다(맞음 {A['V0']['tp']} → {A['1n']['tp']}개).",
          f"설계 3 은 수집 영상에서 초과 등록을 {LA['V0']['extra']} → {LA['3']['extra']}개로 늘렸습니다. 설계 2 는 깊이 지지율을 {LA['V0']['sup']:.2f} → {LA['2']['sup']:.2f} 로 올렸지만 SmallRoom 맞음은 {A['V0']['tp']} → {A['2']['tp']}개로 줄었습니다."])
    d.h3("지도 갱신에 걸린 시간")
    trow = lambda n, o: [n, o["passes"], f"{o['V0_median']:.0f} / {o['V0_max']:.0f}", f"{o['1n_median']:.0f} / {o['1n_max']:.0f}", f"{o['1n+2_median']:.0f} / {o['1n+2_max']:.0f}", f"{o['1n+3_median']:.0f} / {o['1n+3_max']:.0f}", f"{o['all_med']:.0f} / {o['all_max']:.0f}", f"{o['lag_med']:.0f} / {o['lag_max']:.0f}"]
    TH = ["구성", "갱신 횟수", "지금 방식 + 지지율 측정", "설계 1′", "설계 1′+2", "설계 1′+3", "열두 가지 모두", "인식 결과 도착 → 지도 반영"]
    d.p("단위 ms, 중앙값 / 최대. 방식별 값은 그 방식 하나만 돌렸을 때의 시간입니다(영상별 중앙값의 중앙값, 영상별 최대의 최대).")
    d.table(["SmallRoom"] + TH[1:], [trow(NO[c], ON_OB[c]) for c in CS]); d.table(["수집 영상"] + TH[1:], [trow(NO[c], ON_LAB[c]) for c in CS])
    d.ul(["합의만 쓰면(설계 1′) 갱신 한 번이 중앙값 10 ms 안팎이고, 깊이로 다듬는 설계 2 를 넣으면 수백 ms 가 됩니다.",
          "수집 영상에서는 열두 가지를 한꺼번에 계산하느라 갱신 한 번이 중앙값 0.6 ~ 0.9초, 최대 3.7초까지 걸렸습니다. 한 가지 방식만 돌리는 실제 사용에서는 위의 방식별 값이 해당됩니다.",
          "실행 중에 계산한 끝 시점 지도가 같은 기록의 사후 계산과 같다는 것은 SmallRoom 방문 하나에서 확인했습니다(위치 차 0.4 mm 이하)."])
    d.h2("6. 반복 실험: 같은 인식 결과에 지도 생성 방식만 바꾸면")
    d.p("제안 인식(③)을 실시간으로 돌려 얻은 인식 결과 하나를 여섯 가지 지도 생성 방식에 똑같이 넣었습니다. M0~M3 은 실행 중에 계산했고, M4 · M5 는 저장한 같은 인식 결과(마스크와 이름)를 FM-Fusion, DualMap 의 원래 지도 생성 프로그램에 넣었습니다. 전체를 3회 반복했고 괄호와 범위는 3회의 최소~최대입니다.")
    d.h3("6-1. ObChange 방 네 곳 (장면 20개, 정답 객체 %d개)" % RP["ngt"])
    d.table(["지도 생성 방식", "맞음", "잘못 등록", "F1 % 평균 (범위)", "겹침 0.5 이상", "평균 겹침"], RP["OB"], cls="sum")
    d.h3("6-2. YCB-Video 장면 12개 (정답 객체 55개, 정답 자세 있음)")
    d.table(["지도 생성 방식", "맞음", "잘못 등록", "F1 %", "평균 겹침", "ADD-S 중앙값 mm", "위치 오차 중앙값 mm", "위치 오차 90 % mm"], RP["YC"], cls="sum")
    d.h3("6-3. 인식 구성 네 가지, ObChange 방 네 곳 반복 (정답 객체 %d개, 지도는 기존 갱신)" % RP["ngt"])
    d.table(["구성", "실행 횟수", "맞음", "잘못 등록", "F1 % 평균 (범위)", "SmallRoom F1", "Office F1", "LivingArea F1", "Kitchen F1"], RP["NB"], cls="sum")
    d.h3("6-4. 관문 대 총점 판정, YCB-Video 구성마다 3회 (정답 객체 55개)")
    d.table(["구성, 지도 생성", "맞음", "잘못 등록", "F1 % 평균 (범위)", "ADD-S 중앙값 mm", "위치 오차 중앙값 mm"], RP["NY"], cls="sum")
    d.h3("6-5. YCB-M 지도 생성 방식 3회 (정답 객체 %d개, 인식은 ③)" % RP["ymgt"])
    d.table(["지도 생성 방식", "맞음", "잘못 등록", "F1 % 평균 (범위)", "ADD-S 중앙값 mm", "위치 오차 중앙값 mm", "위치 오차 90 % mm"], RP["YM"], cls="sum")
    d.ul(["FM-Fusion · DualMap 의 지도 생성(M4 · M5)보다 지금 방식 계열(M1 · M2)의 F1 이 ObChange 3회 모두 같거나 높았습니다. YCB-Video 에서는 FM-Fusion 이 한 번 같은 값까지 올라왔습니다.",
          "단순 누적(M0)보다 F1 이 높다고는 말할 수 없습니다. 지금 방식은 잘못 등록을 줄이지만 ObChange 에서 맞음도 함께 줄었습니다.",
          "합의로 포즈를 정하면(M2) 맞음 · 잘못 등록은 M1 과 같고 위치가 정확해집니다. 단순 누적(M0)의 위치 오차도 M2 와 같은 수준이어서, 이 효과는 M1 대비로만 말할 수 있습니다.",
          "깊이로 다듬기(M3)는 효과가 없었습니다. YCB-Video 에서 위치 오차 90 % 값이 오히려 커졌습니다.",
          "M4 · M5 는 6D 자세를 저장하지 않아 자세 오차가 없습니다. 자세가 있는 M0~M3 은 맞게 등록된 객체가 모두 ADD-S 기준(지름의 10 %)을 통과했습니다.",
          "방 네 곳을 합치면 제안 관문(③)은 총점 판정(⑤)과 맞음이 같은 수준이고 잘못 등록이 37~42건에서 23~25건으로 줄었습니다. 방별로는 SmallRoom · Office 에서 관문 판정(② · ③)이, Kitchen 에서는 총점 판정(① · ⑤)이 높았고, LivingArea 는 ⑤ 만 낮았습니다.",
          "후보 생성만 다른 ② 와 ③ 은 방 네 곳 F1 평균이 같습니다(62.5 대 62.3). 문구 상자가 화면 전체 분할보다 낫다고는 말할 수 없습니다.",
          "YCB-Video 에서는 총점 판정(① · ⑤)이 관문 판정(② · ③)보다 4~6개쯤 더 맞혔고 잘못 등록은 두 배쯤 많았습니다. F1 평균은 화면 전체 분할에서는 관문이(76.0 → 80.6), 문구 상자에서는 총점 판정이 높았습니다(83.4 대 81.5). 탁자 위 장면에서는 관문이 맞는 객체도 걸러 냅니다.",
          "YCB-M 에서도 지도 생성 방식의 순서는 같습니다. 지금 방식 계열은 FM-Fusion · DualMap 의 지도 생성보다 F1 이 같거나 높고, 단순 누적보다는 잘못 등록이 적지만 맞음이 줄어 F1 이 조금 낮습니다. 합의의 위치 개선은 YCB-M 에서 1 mm 이내로 작았습니다.",
          "6-3 · 6-4 는 구성마다 3회입니다(③ 은 6-1 의 M1, ① · ② 의 ObChange 3회째와 YCB-Video 3회는 2026-10-13 실행). 결과 <code>objpose/output/mapgen_261012/</code>, 실행 <code>objpose/tools/mapgen/</code>"])
    d.h2("7. 제안 방법(③)이 만든 객체 지도")
    d.p("복원한 장면 위에 지도에 저장된 6D 자세대로 CAD 점과 좌표축을 놓은 그림입니다. SmallRoom 은 이름표에 맞음 / 잘못 등록, 수집 영상은 깊이 지지율을 함께 적었습니다.")
    for f in sorted(MED.glob("pose_smallroom_s*_c1.jpg")) + sorted(MED.glob("pose_26*_c1.jpg")): d.fig(f, ("ObChange SmallRoom 방문 " + f.stem.split("_s")[1][0] if "smallroom" in f.name else "직접 수집한 영상 " + f.stem[5:-3]) + " — 제안 방법(③), 지금 방식의 지도")
    d.h2("8. 한계와 재현")
    d.ul(["1~5절의 실행은 구성마다 한 번입니다. 어젯밤 같은 구성(⑤)의 SmallRoom F1 은 50.9 %, 오늘은 " + f1(T[(6, 'V0')]['F']) + " % 였습니다.",
          "ObChange 는 SmallRoom 한 방만 돌렸습니다. 수집 영상은 정답이 없어 등록 종류 · 초과 등록 · 깊이 지지율로만 봅니다.",
          "⑥ · ⑦ 의 수집 영상 결과는 카메라 위치를 낸 SLAM 실행이 ①~⑤ 와 다릅니다.",
          "실행: <code>objpose/tools/lab/run_final5.sh</code> (<code>run_live_rt.sh</code>, <code>objpose/tools/obchange/run_rt_obchange.sh</code>), 인식기 <code>objpose/tools/lab/live_six.py</code>(LIVE_SIX 2 · 9 · 8 · 7 · 6 = ① ~ ⑤), 실행 중 지도 설계 <code>objpose/pc/online_designs.py</code>",
          "채점: <code>objpose/tools/obchange/score_designs.py</code>, <code>score_base_same.py</code>. 결과 <code>objpose/output/rt_obchange_261011/</code>, <code>objpose/output/live_rt_261011/</code>"])
    ref = REF.read_text(); head = ref[: ref.index("<div class=\"wrap\">")].replace("<title>YCB 조합 실험 보고</title>", "<title>확정 구성 실시간 결과</title>").replace("</style>", "figure{margin:6px 0 4px}figure img{max-width:100%;height:auto;border:1px solid var(--line);border-radius:4px;display:block;background:#fff}figcaption{margin-top:4px;margin-bottom:10px}</style>")
    hdr = ('<header><div class="kicker">확정 구성 일곱 가지 · 실시간 객체 지도 · 2026-10-11</div><h1>확정 구성 일곱 가지의 실시간 객체 지도 결과</h1>'
           '<p class="lead">객체 자세 추정 구성 다섯 가지와 FM-Fusion, DualMap 을 직접 수집한 영상 8개와 ObChange SmallRoom 에 실시간으로 돌려 최종 객체 지도를 비교했습니다. 지도 생성 설계는 실행 중에 계산했습니다.</p></header>')
    body = hdr + "\n" + "\n".join(x for x, _ in d.parts)
    (HERE / "final7_report.html").write_text(head + '<div class="wrap">\n' + body + "\n</div>\n")
    (HERE / "final7_report_local.html").write_text('<!doctype html>\n<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"></head><body style="margin:0">\n' + head + '<div class="wrap">\n' + body + "\n</div>\n</body></html>\n")
    nref = (REF.parent / "ycb_combo_report_notion.html").read_text(); nhead = nref[: nref.index("<body>") + 6].replace("YCB 조합 실험 보고", "확정 구성 실시간 결과")
    nb = ("<p>확정 구성 일곱 가지 · 2026-10-11</p>" + re.sub(r'<header>.*?</div>', "", hdr, flags=re.S).replace("</header>", "").replace(' class="lead"', "")) + "\n" + "\n".join(y for _, y in d.parts)
    (HERE / "final7_report_notion.html").write_text(nhead + "\n\n" + nb + "\n</body>\n</html>\n")
    print("html ok | best", bestF, {NO[c]: (round(T[(c, 'V0')]['F'], 1), LT[(c, 'V0')]['kinds'], LT[(c, 'V0')]['extra']) for c in CS}, {m: round(BT[m]["F"], 1) for m in BT}, LB)
    if a.pptx: pptx(E)


def pptx(E):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    for x in ("docs/ppt/dataset_exp_261005", "docs/ppt/arch_261005", "tools/analysis"): sys.path.insert(0, str(REPO / "objpose" / x))
    from build import slide, lines, LM, CW, W
    sys.path.append(str(REPO / "objpose/docs/ppt/fair_results_261005"))
    from replace_tables import title_of
    from append_rooms import grid
    T, A, BT, LT, LA, LB, ON_OB, ON_LAB, CMP, bestF = (E[k] for k in ("T", "A", "BT", "LT", "LA", "LB", "ON_OB", "ON_LAB", "CMP", "bestF"))
    OUT = Path.home() / "콜로키움" / "전체_실험결과.pptx"; PRE = "종합"; prs = Presentation(str(OUT)); ids = prs.slides._sldIdLst; ngt = T[(8, "V0")]["gt"]
    for i in sorted((i for i, s in enumerate(prs.slides) if title_of(s).startswith(PRE + ":")), reverse=True): prs.part.drop_rel(ids[i].rId); ids.remove(ids[i])
    SN = {2: "① 원본", 9: "② 관문 비교", 8: "③ 제안 방법", 7: "④ 완화 비교", 6: "⑤ 총점 효과"}; hi = lambda rows: (lambda i, j: "e2efda" if rows[i][0].startswith("③") else "ffffff")
    # 1 구성
    s = slide(prs, PRE + ": 비교한 구성 일곱 가지", "자세 추정 구성 다섯 가지와 다른 객체 지도 방법 두 가지를 실시간으로 돌려 최종 객체 지도를 비교했다.", "")
    R = [["구성", "후보 생성", "마스크", "물체 판정"]] + [[SN[c], *SPEC[c]] for c in [2, 9, 8, 7, 6]] + [["⑥ FM-Fusion", "GroundingDINO 문구 상자", "SAM", "검출기의 이름"], ["⑦ DualMap", "YOLO-World 문구 상자", "MobileSAM / FastSAM", "검출기의 이름 + CLIP"]]
    y = grid(s, R, [2.1, 3.5, 2.6, 3.93], 1.75, 0.4, 0.46, 12, fill=hi(R), left=(1, 2, 3))
    lines(s, LM, y + 0.25, CW, 2.2, ["①~⑤ 는 특징 ViT-S, 자세 추정, 자세 검증을 모두 켰다. ③ · ④ · ⑤ 는 후보 상자와 마스크가 같고 물체 판정만 다르다.",
                                     f"데이터: 직접 수집한 영상 8개(노트북의 SLAM 이 실시간으로 낸 카메라 위치)와 ObChange SmallRoom 방문 5회(정답 물체 {ngt}개, 10 fps 실시간 재생).",
                                     "⑥ · ⑦ 은 물체를 점 묶음으로 저장하는 자체 지도 생성 방식을 그대로 썼다."], sz=13, gap=6)
    # 2 SmallRoom
    s = slide(prs, PRE + ": ObChange SmallRoom 결과", f"제안 방법(③)이 F1 {f1(T[(8, 'V0')]['F'])} % 로 가장 높다. 차이는 잘못 등록한 객체 수에서 난다.", "")
    row = lambda n, t: [n, f"{t['tp']} / {ngt}", t["fp"], f1(t["P"]), f1(t["R"]), f1(t["F"]), t["tp50"]]
    R = [["구성", "맞음", "잘못 등록", "정밀도 %", "재현율 %", "F1 %", "겹침 0.5 이상"]] + [row(SN[c], T[(c, "V0")]) for c in [2, 9, 8, 7, 6]] + [row("⑥ FM-Fusion", BT["FM-Fusion"]), row("⑦ DualMap", BT["DualMap"])]
    y = grid(s, R, [2.6, 1.6, 1.6, 1.6, 1.6, 1.5, 1.63], 1.75, 0.4, 0.44, 12.5, fill=hi(R))
    lines(s, LM, y + 0.25, CW, 2.3, [f"총점을 관문으로 바꾸면 잘못 등록이 줄어든다: 문구 상자에서 {T[(6, 'V0')]['fp']} → {T[(8, 'V0')]['fp']}개(⑤ → ③), 화면 전체 분할에서 {T[(2, 'V0')]['fp']} → {T[(9, 'V0')]['fp']}개(① → ②).",
                                     f"관문을 완화하면(④) 잘못 등록이 {T[(8, 'V0')]['fp']} → {T[(7, 'V0')]['fp']}개로 다시 늘어난다.",
                                     f"FM-Fusion 과 DualMap 은 맞은 수는 비슷하지만 이름이 붙은 잘못된 객체가 {BT['FM-Fusion']['fp']}, {BT['DualMap']['fp']}개다.",
                                     "맞음: 정답과 같은 이름의 객체가 0.25 이상 겹침. 실행은 구성마다 한 번이다."], sz=13, gap=6)
    # 3 수집 영상
    s = slide(prs, PRE + ": 직접 수집한 영상 8개 결과", f"제안 방법(③)은 물체 64개를 모두 등록하면서 같은 물체를 두 번 등록한 경우가 {LT[(8, 'V0')]['extra']}건이다.", "")
    R = [["구성", "등록된 종류 (최대 64)", "초과 등록", "깊이 지지율"]] + [[SN[c], LT[(c, "V0")]["kinds"], LT[(c, "V0")]["extra"], f"{LT[(c, 'V0')]['sup']:.2f}"] for c in [2, 9, 8, 7, 6]] + [["⑥ FM-Fusion", LB["FM-Fusion"]["kinds"], LB["FM-Fusion"]["extra"], "해당 없음"], ["⑦ DualMap", LB["DualMap"]["kinds"], LB["DualMap"]["extra"], "해당 없음"]]
    y = grid(s, R, [3.2, 3.2, 2.9, 2.83], 1.75, 0.4, 0.44, 12.5, fill=hi(R))
    lines(s, LM, y + 0.25, CW, 2.3, ["영상마다 물체 8종이 하나씩 놓여 있다. 정답 자세가 없어 등록된 종류와 초과 등록(같은 종류를 한 번 넘게 등록)으로 비교했다.",
                                     f"총점 판정은 초과 등록이 많다(① {LT[(2, 'V0')]['extra']}건, ⑤ {LT[(6, 'V0')]['extra']}건). 관문 판정은 {LT[(9, 'V0')]['extra']}건(②), {LT[(8, 'V0')]['extra']}건(③)이다.",
                                     "⑥ · ⑦ 은 카메라 위치를 낸 SLAM 실행이 ①~⑤ 와 다르다(같은 ORB-SLAM3, 별도 실행)."], sz=13, gap=6)
    # 3-2 한 장 평가와 지도 평가
    SV, SM = STILL["YCB-Video 900장"], STILL["YCB-M 925장"]; g = lambda D, c, i, fmt: (fmt.format(D[c][i]) if c in D else "-")
    s = slide(prs, PRE + ": 사진 한 장 평가와 객체 지도 평가", "사진 한 장씩 평가하면 총점 판정이 더 많이 찾는다. 객체 지도로 평가하면 제안 관문이 잘못 등록을 줄여 앞선다.", "")
    R = [["구성", "한 장 평가 찾은 비율 %", "한 장 평가 오답 / 장", "지도 평가 맞음", "지도 평가 잘못 등록", "지도 평가 F1 %", "수집 영상 초과 등록"]] + [[SN[c], g(SV, c, 0, "{:.1f}"), g(SV, c, 2, "{:.3f}"), f"{T[(c, 'V0')]['tp']} / {ngt}", T[(c, "V0")]["fp"], f1(T[(c, "V0")]["F"]), LT[(c, "V0")]["extra"]] for c in [2, 9, 8, 7, 6]]
    y = grid(s, R, [2.6, 1.7, 1.6, 1.5, 1.6, 1.5, 1.63], 1.75, 0.62, 0.46, 12.5, fill=hi(R))
    lines(s, LM, y + 0.25, CW, 2.5, [f"한 장 평가(YCB-Video 900장, 물체 21종을 모두 묻는 조건): 총점 판정 ⑤ 는 {SV[6][0]} % 를 찾고 제안 관문 ③ 은 {SV[8][0]} % 를 찾았다. 오답은 {SV[6][2]:.3f} 대 {SV[8][2]:.3f} / 장으로 같은 수준이다. YCB-M 925장도 같은 방향이다({SM[6][0]} 대 {SM[8][0]} %).",
                                     f"지도 평가(ObChange SmallRoom, 실시간): 맞은 수는 {T[(6, 'V0')]['tp']} 대 {T[(8, 'V0')]['tp']} 로 같은 수준이고, 잘못 등록은 {T[(6, 'V0')]['fp']} 대 {T[(8, 'V0')]['fp']} 다.",
                                     "두 평가는 세는 대상과 데이터가 다르다. 한 장 평가는 프레임마다의 자세를, 지도 평가는 영상이 끝난 뒤 지도에 남은 객체를 센다.",
                                     "한 장 평가의 ③ 은 원래 문구, ① · ⑤ 는 새 문구를 썼다. ② 는 한 장 평가 전체 실행이 없다."], sz=13, gap=6)
    # 4 지도 방식
    s = slide(prs, PRE + ": 객체 지도 생성 방식 비교", "영상을 실시간으로 넣으면서 지도까지 그 자리에서 만든 결과다. 합의로 합치면 자세는 정확해지지만 맞은 수가 줄어 지금 방식을 넘지 못했다.", "")
    vs = ["V0", "1n", "2", "3", "1n+2", "1n+3"]; nm = {"V0": "지금 방식 (위치 평균, 방향 다수결)", "1n": "합의로 합치기", "2": "깊이로 확인하고 다듬기", "3": "약한 증거 누적", "1n+2": "합의 + 깊이", "1n+3": "합의 + 약한 증거"}
    R = [["지도 생성 방식", "맞음", "잘못 등록", "F1 %", "겹침 0.5 이상", "수집 영상 초과 등록", "갱신 한 번 (ms)"]] + [[nm[v], A[v]["tp"], A[v]["fp"], f1(A[v]["F"]), A[v]["tp50"], LA[v]["extra"], (f"{np.median([ON_LAB[c][v + '_median'] for c in ON_LAB]):.0f}" if v + "_median" in ON_LAB[8] else "-")] for v in vs]
    y = grid(s, R, [4.0, 1.2, 1.4, 1.2, 1.6, 1.4, 1.33], 1.75, 0.4, 0.44, 12, left=(0,))
    lines(s, LM, y + 0.25, CW, 2.3, ["다섯 구성의 합이다(맞음 · 잘못 등록 · F1 · 겹침은 ObChange SmallRoom, 정답 물체 " + str(A["V0"]["gt"]) + "개). 갱신 시간은 수집 영상에서 그 방식 하나만 돌렸을 때의 중앙값이다.",
                                     f"합의로 합치기는 겹침 0.5 이상인 객체를 {A['V0']['tp50']} → {A['1n']['tp50']}개로 늘렸지만 맞음이 {A['V0']['tp']} → {A['1n']['tp']}개로 줄었다. 무리를 잘못 고르면 평균보다 더 벗어난다.",
                                     f"약한 증거 누적은 수집 영상의 초과 등록을 {LA['V0']['extra']} → {LA['3']['extra']}개로 늘렸다."], sz=13, gap=6)
    # 4-2 반복 실험
    RP = E["RP"]; hm = lambda rows: (lambda i, j: "e2efda" if rows[i][0].startswith(("M1", "M2", "③")) else "ffffff")
    s = slide(prs, PRE + ": 같은 인식 결과로 지도 생성 방식만 바꾼 반복 실험 (방 크기)", "제안 인식 결과 하나를 여섯 가지 지도 생성 방식에 넣어 3회 반복했다. 지금 방식 계열이 다른 두 방법의 지도 생성보다 F1 이 높다.", "")
    R = [["지도 생성 방식", "맞음", "잘못 등록", "F1 % 평균 (범위)", "겹침 0.5 이상", "평균 겹침"]] + RP["OB"]
    y = grid(s, R, [4.2, 1.5, 1.5, 2.3, 1.4, 1.23], 1.75, 0.4, 0.44, 12, fill=hm(R), left=(0,))
    lines(s, LM, y + 0.25, CW, 2.3, [f"ObChange 방 네 곳, 장면 20개, 정답 객체 {RP['ngt']}개. 범위는 3회의 최소~최대다. M4 · M5 는 같은 마스크와 이름을 원래 프로그램에 넣었다.",
                                     "지금 방식(M1)은 단순 누적(M0)보다 잘못 등록이 적지만 맞음도 줄어 F1 은 차이가 없다.",
                                     "합의로 포즈를 정하면(M2) 맞음 · 잘못 등록은 그대로이고 겹침 0.5 이상인 객체가 늘어난다. 깊이로 다듬기(M3)는 효과가 없다."], sz=13, gap=6)
    s = slide(prs, PRE + ": 지도에 저장된 6D 자세의 오차 (YCB-Video 반복 실험)", f"합의로 포즈를 정하면 위치 오차 중앙값이 {RP['YC'][1][6]} mm 에서 {RP['YC'][2][6]} mm 로 줄어든다. 등록 수는 변하지 않는다.", "")
    R = [["지도 생성 방식", "맞음", "잘못 등록", "F1 %", "ADD-S 중앙값 mm", "위치 오차 중앙값 mm", "위치 오차 90 % mm"]] + [[r[0], r[1], r[2], r[3], r[5], r[6], r[7]] for r in RP["YC"]]
    y = grid(s, R, [4.0, 1.2, 1.3, 1.5, 1.4, 1.4, 1.33], 1.75, 0.62, 0.44, 12, fill=hm(R), left=(0,))
    lines(s, LM, y + 0.25, CW, 2.3, ["YCB-Video 장면 12개, 정답 객체 55개, 30 fps 실시간 재생, 3회 반복의 최소~최대다.",
                                     "FM-Fusion 과 DualMap 의 지도 생성은 객체를 점 묶음으로 저장하므로 자세 오차를 낼 수 없다.",
                                     "자세가 있는 M0~M3 은 맞게 등록된 객체가 모두 ADD-S 기준(지름의 10 %)을 통과했다. 단순 누적(M0)의 위치 오차도 합의(M2)와 같은 수준이다."], sz=13, gap=6)
    NB, NY, YM = RP["NB"], RP["NY"], RP["YM"]
    s = slide(prs, PRE + ": 인식 구성 반복 실험 (방 네 곳)", f"제안 관문은 총점 판정과 맞은 수가 같은 수준이고 잘못 등록이 {NB[3][3]}건에서 {NB[2][3]}건으로 줄어든다. 방에 따라 순서는 달라진다.", "")
    R = [["구성", "실행 횟수", "맞음", "잘못 등록", "F1 % 평균 (범위)", "SmallRoom", "Office", "LivingArea", "Kitchen"]] + NB
    y = grid(s, R, [3.3, 0.85, 0.95, 1.0, 1.95, 1.1, 0.8, 1.2, 0.98], 1.75, 0.62, 0.46, 11.5, fill=hm(R), left=(0,))
    lines(s, LM, y + 0.25, CW, 2.3, [f"ObChange 방 네 곳, 장면 20개, 정답 객체 {RP['ngt']}개, 10 fps 실시간 재생. 오른쪽 네 칸은 방별 F1 평균이다.",
                                     "SmallRoom · Office 는 관문 판정 두 구성(② · ③)이, Kitchen 은 총점 판정 두 구성(① · ⑤)이 높다. LivingArea 는 ⑤ 만 낮다.",
                                     f"후보 생성만 다른 ② 와 ③ 은 F1 평균이 {NB[1][4].split(' ')[0]} 대 {NB[2][4].split(' ')[0]} 로 비슷하다. 차이를 만드는 것은 물체 판정(관문)이다."], sz=13, gap=6)
    s = slide(prs, PRE + ": 관문 대 총점 판정 (YCB-Video)", f"관문을 쓰면 잘못 등록이 {NY[0][2]} → {NY[2][2]}건, {NY[6][2]} → {NY[4][2]}건으로 줄고, 맞은 수도 {NY[0][1]} → {NY[2][1]}개, {NY[6][1]} → {NY[4][1]}개로 줄어든다.", "")
    R = [["구성, 지도 생성", "맞음", "잘못 등록", "F1 % 평균 (범위)", "ADD-S 중앙값 mm", "위치 오차 중앙값 mm"]] + NY
    y = grid(s, R, [3.6, 1.5, 1.6, 2.4, 1.5, 1.53], 1.75, 0.5, 0.4, 12, fill=hm(R), left=(0,))
    lines(s, LM, y + 0.25, CW, 2.3, ["YCB-Video 장면 12개, 정답 객체 55개, 30 fps 실시간 재생, 구성마다 3회 반복이다.",
                                     "물체가 가깝고 크게 보이는 장면에서는 관문이 맞는 객체도 걸러 낸다. F1 평균은 화면 전체 분할에서는 관문이, 문구 상자에서는 총점 판정이 높았다.",
                                     "합의로 포즈를 정하면 ② 를 뺀 세 구성에서 위치 오차 중앙값이 절반 수준으로 줄어든다. ② 는 기존 갱신에서도 3.5~4.0 mm 였다."], sz=13, gap=6)
    s = slide(prs, PRE + ": 지도 생성 방식 반복 실험 (YCB-M)", "YCB-M 에서도 지금 방식 계열은 다른 두 방법의 지도 생성보다 F1 이 같거나 높고, 단순 누적보다는 조금 낮다.", "")
    R = [["지도 생성 방식", "맞음", "잘못 등록", "F1 % 평균 (범위)", "ADD-S 중앙값 mm", "위치 오차 중앙값 mm", "위치 오차 90 % mm"]] + YM
    y = grid(s, R, [3.6, 1.3, 1.3, 2.2, 1.3, 1.3, 1.13], 1.75, 0.62, 0.44, 12, fill=hm(R), left=(0,))
    lines(s, LM, y + 0.25, CW, 2.3, [f"YCB-M 정답 객체 {RP['ymgt']}개, 실시간 재생, 3회 반복의 최소~최대다. 인식은 제안 방법으로 같다.",
                                     "지금 방식(M1)은 단순 누적(M0)보다 잘못 등록이 적지만 맞음이 줄어 F1 은 2점쯤 낮다.",
                                     "합의(M2)의 위치 개선은 1 mm 이내로 YCB-Video 보다 작다. FM-Fusion 의 지도 생성은 한 번 같은 수준까지 올라왔다."], sz=13, gap=6)
    # 5 지도 그림
    pics = [MED / "pose_smallroom_s6_c1.jpg", MED / "pose_260901_cbnu_longcircle_c1.jpg"]; pics = [p for p in pics if p.exists()]
    if pics:
        s = slide(prs, PRE + ": 제안 방법이 만든 객체 지도", "복원한 장면 위에 지도에 저장된 6D 자세대로 CAD 모형과 좌표축을 놓았다.", "")
        w = (W - 2 * 0.5 - 0.3) / len(pics)
        for i, p in enumerate(pics):
            iw, ih = Image.open(p).size; h = min(4.6, w * ih / iw); ww = h * iw / ih; s.shapes.add_picture(str(p), Inches(0.5 + i * (w + 0.3) + (w - ww) / 2), Inches(1.8), Inches(ww), Inches(h))
        lines(s, LM, 6.55, CW, 0.6, ["왼쪽: ObChange SmallRoom(이름표에 맞음 / 잘못 등록). 오른쪽: 직접 수집한 영상(이름표에 깊이 지지율)."], sz=12, gap=4)
    prs.part.rename_slide_parts([x.rId for x in ids]); prs.save(str(OUT)); print("pptx slides", len(ids), "종합", sum(1 for s in prs.slides if title_of(s).startswith(PRE + ":")))


if __name__ == "__main__":
    main()
